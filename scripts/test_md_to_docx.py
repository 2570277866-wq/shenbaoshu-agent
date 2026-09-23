# -*- coding: utf-8 -*-
"""
md_to_docx 单测。

python-docx 是可选依赖：系统 python3 没装时整组跳过（跑全量不红），
装了（agent_service 的 .venv）才真正执行。因此 agent_service 的测试
会在自己的 venv 里把本文件再跑一遍。
"""

import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import md_to_docx  # noqa: E402

DOCX = md_to_docx._HAS_DOCX


@unittest.skipUnless(DOCX, "需要 python-docx（agent_service venv 里跑）")
class TestConvert(unittest.TestCase):
    def texts(self, doc):
        return [p.text for p in doc.paragraphs]

    def test_headings(self):
        doc = md_to_docx.convert("# 标题\n## 一节\n### 小节\n正文。")
        self.assertEqual(self.texts(doc)[:4],
                         ["标题", "一节", "小节", "正文。"])
        self.assertEqual(doc.paragraphs[0].style.name, "Title")
        self.assertEqual(doc.paragraphs[1].style.name, "Heading 1")
        self.assertEqual(doc.paragraphs[2].style.name, "Heading 2")

    def test_bold_inline(self):
        doc = md_to_docx.convert("正文 **加粗** 尾巴。")
        para = doc.paragraphs[0]
        self.assertEqual(para.text, "正文 加粗 尾巴。")
        self.assertTrue(para.runs[1].font.bold)
        self.assertFalse(para.runs[0].font.bold)

    def test_table(self):
        doc = md_to_docx.convert(
            "| 科目 | 金额 |\n|---|---|\n| 设备费 | 45 万 |\n| **合计** | 200 万 |")
        self.assertEqual(len(doc.tables), 1)
        table = doc.tables[0]
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(table.cell(0, 0).text, "科目")
        self.assertEqual(table.cell(2, 0).text, "合计")
        self.assertTrue(table.cell(0, 0).paragraphs[0].runs[0].font.bold)

    def test_placeholder_preserved(self):
        doc = md_to_docx.convert("周期【待补充：项目周期】。")
        self.assertIn("【待补充：项目周期】", doc.paragraphs[0].text)

    def test_lists(self):
        doc = md_to_docx.convert("- 甲\n- 乙\n1. 一\n2. 二")
        styles = [p.style.name for p in doc.paragraphs]
        self.assertEqual(styles[:2], ["List Bullet", "List Bullet"])
        self.assertEqual(styles[2:], ["List Number", "List Number"])

    def test_empty_input(self):
        doc = md_to_docx.convert("")
        self.assertEqual(self.texts(doc), [])

    def test_east_asia_font_set(self):
        doc = md_to_docx.convert("正文。")
        run = doc.paragraphs[0].runs[0]
        self.assertEqual(run._element.rPr.rFonts.get(md_to_docx.qn("w:eastAsia")),
                         "宋体")

    def test_title_override(self):
        doc = md_to_docx.convert("正文。", title="项目申报书")
        self.assertEqual(doc.paragraphs[0].text, "项目申报书")
        self.assertEqual(doc.paragraphs[0].style.name, "Title")

    def test_save_and_reopen(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.docx")
            md_to_docx.save(md_to_docx.convert("# 标题\n正文。"), path)
            from docx import Document
            reopened = Document(path)
            self.assertEqual([p.text for p in reopened.paragraphs][:2],
                             ["标题", "正文。"])

    def test_cli(self):
        with tempfile.TemporaryDirectory() as tmp:
            md = os.path.join(tmp, "doc.md")
            out = os.path.join(tmp, "doc.docx")
            with open(md, "w", encoding="utf-8") as fh:
                fh.write("# 标题\n正文。")
            self.assertEqual(md_to_docx.main([md, out]), 0)
            self.assertTrue(os.path.exists(out))
            with open(out, "rb") as fh:
                self.assertEqual(fh.read(2), b"PK")


@unittest.skipIf(DOCX, "python-docx 已装，错误只在未装时测")
class TestMissingDep(unittest.TestCase):
    def test_convert_raises_clear_error(self):
        with self.assertRaises(RuntimeError) as ctx:
            md_to_docx.convert("正文")
        self.assertIn("pip install python-docx", str(ctx.exception))


if __name__ == "__main__":
    unittest.main()
